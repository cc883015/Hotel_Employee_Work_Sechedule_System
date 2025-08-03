# 完整代码：单文件合并 + 窗口适配 + 日历查看详情 + 导出命名 + 管理员工 + 排班功能保留 + UI美化 + 月份修正 + 背景图支持 + 日历美化增强 + 按钮Emoji + 功能回归 + 按人名排序搜索优化 + Daily视图增加星期显示 + 窗口宽度可缩放 + 删除非日期红字 + Cleaner工资可设置 + 排班与导出功能还原 + 导出工资容错 + 导出路径选择功能 + 日期数字改表情 + 排班控件完整恢复
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
from tkcalendar import Calendar, DateEntry
import sqlite3
from datetime import datetime
import pandas as pd
from PIL import Image, ImageTk
import calendar as calmod
import openpyxl

root = tk.Tk()
root.title("Cliff House Employee Scheduler")
root.geometry("1300x800")
root.resizable(True, True)

CLEANER_PAY = tk.IntVar(value=120)
RECEPTIONIST_FULL = 150
RECEPTIONIST_HALF = 75

conn = sqlite3.connect("work_schedule.db")
cur = conn.cursor()
cur.execute('''CREATE TABLE IF NOT EXISTS schedule (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT,
    name TEXT,
    role TEXT,
    shift TEXT
)''')
conn.commit()

employees = ["Ashley", "News", "Ethan", "Charles", "Kenny", "Sami"]

# 设置背景图
try:
    bg_img = Image.open("background.png")
    bg_img = bg_img.resize((1300, 800))
    bg_photo = ImageTk.PhotoImage(bg_img)
    background_label = tk.Label(root, image=bg_photo)
    background_label.place(relx=0, rely=0, relwidth=1, relheight=1)
except:
    root.configure(bg="#f4f4f4")

bold_font = ("Arial", 13, "bold")

# ========== 顶部标题 ==========
top_frame = tk.Frame(root, bg="white")
top_frame.grid(row=0, column=0, columnspan=4, pady=(10, 0))

try:
    logo_img = Image.open("logo.png")
    logo_img = logo_img.resize((40, 40))
    logo_photo = ImageTk.PhotoImage(logo_img)
    logo_label = tk.Label(top_frame, image=logo_photo, bg="white")
    logo_label.pack(side="left", padx=10)
except:
    logo_label = tk.Label(top_frame, text="🏠", font=("Arial", 24), bg="white")
    logo_label.pack(side="left", padx=10)

app_title = tk.Label(top_frame, text="Cliff House Employee Work System", font=("Arial", 22, "bold"), bg="white")
app_title.pack(side="left")

# ========== 日历区域 ==========
frame_calendar = tk.LabelFrame(root, text="📅 Choose Date", padx=10, pady=5, font=bold_font, bg="white")
frame_calendar.grid(row=1, column=0, padx=10, pady=5, sticky='nw')

class EmojiCalendar(Calendar):
    def _display_calendar(self):
        super()._display_calendar()
        emoji = ['🐶', '🐱', '🐭', '🐹', '🐰', '🦊']
        if hasattr(self, '_calendar') and hasattr(self, '_items'):
            for i, e in enumerate(emoji):
                if i + 1 < len(self._items):
                    self._calendar.item(self._items[i + 1], text=e)

cal = EmojiCalendar(frame_calendar,
               selectmode='day',
               year=datetime.now().year,
               month=datetime.now().month,
               day=datetime.now().day,
               date_pattern='yyyy-mm-dd',
               font=("Arial", 18),
               selectforeground='white',
               selectbackground='blue',
               headersforeground='red',
               headersbackground='white',
               foreground='black',
               background='white',
               weekendbackground='lightgrey',
               bordercolor='black',
               othermonthwebackground='white',
               othermonthforeground='gray')
cal.pack(padx=5, pady=5, ipadx=45, ipady=20)
cal.calevent_remove('all')

# ========== 查看按钮 ==========
def view_daily_schedule():
    date_obj = cal.selection_get()
    date = date_obj.strftime("%Y-%m-%d")
    weekday = calmod.day_name[date_obj.weekday()]
    cur.execute("SELECT name, role, shift FROM schedule WHERE date=?", (date,))
    rows = cur.fetchall()
    if not rows:
        messagebox.showinfo("No Schedule", f"No one is scheduled on {date} ({weekday}).")
    else:
        result = f"Work on {date} ({weekday}):\n"
        for name, role, shift in rows:
            result += f"- {name} ({role}) — {shift} Day\n"
        messagebox.showinfo("Daily Schedule", result)

btn_view_day = tk.Button(frame_calendar, text="🔍 View Daily", font=bold_font, command=view_daily_schedule)
btn_view_day.pack(pady=5)

# 删除记录按钮
btn_delete_record = tk.Button(frame_calendar, text="🗑 删除记录", font=bold_font, bg="#e74c3c", fg="white", command=delete_record)
btn_delete_record.pack(pady=5)

# ========== 导出函数定义（提前） ==========
def export_schedule(start, end):
    cur.execute("SELECT * FROM schedule WHERE date BETWEEN ? AND ?", (start, end))
    rows = cur.fetchall()
    if not rows:
        messagebox.showinfo("No Data", "No schedule in this period.")
        return
    
    # 准备详细数据
    data = []
    for row in rows:
        pay = 0
        if row[3] == "cleaner":
            pay = CLEANER_PAY.get()
        elif row[3] == "receptionist":
            pay = RECEPTIONIST_FULL if row[4] == "Full" else RECEPTIONIST_HALF
        data.append([row[1], row[2], row[3], row[4], pay])
    
    # 创建详细排班表
    df_detail = pd.DataFrame(data, columns=["Date", "Name", "Role", "Shift", "Pay"])
    
    # 按员工统计 - 半天算0.5天
    employee_stats = []
    for name in df_detail['Name'].unique():
        employee_data = df_detail[df_detail['Name'] == name]
        total_pay = employee_data['Pay'].sum()
        work_days = 0
        for _, row in employee_data.iterrows():
            if row['Shift'] == 'Half':
                work_days += 0.5
            else:
                work_days += 1
        employee_stats.append([name, total_pay, work_days])
    
    employee_stats = pd.DataFrame(employee_stats, columns=['员工姓名', '总工资', '工作天数'])
    
    # 保存到Excel文件
    default_name = f"Schedule_{start}_to_{end}.xlsx"
    file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile=default_name, filetypes=[("Excel files", "*.xlsx")])
    
    if file_path:
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            # 详细排班表
            df_detail.to_excel(writer, sheet_name='详细排班表', index=False)
            
            # 员工统计表
            employee_stats.to_excel(writer, sheet_name='员工统计', index=False)
        
        messagebox.showinfo("导出成功", f"排班表已保存为 {file_path}\n包含以下表格：\n• 详细排班表\n• 员工统计（总工资和工作天数）")

# ========== 排班控件完整写入 ==========
def build_scheduler_controls(root, employees, cur, conn, CLEANER_PAY, RECEPTIONIST_FULL, RECEPTIONIST_HALF, export_schedule):
    frame = tk.Frame(root, bg="white")
    frame.grid(row=1, column=1, padx=5, pady=5, sticky='nw')

    # 工资设置
    tk.Label(frame, text="💰 Set Cleaner Pay", font=bold_font, bg="white").grid(row=0, column=0, sticky='w')
    tk.Entry(frame, textvariable=CLEANER_PAY, width=10).grid(row=0, column=1, sticky='w')

    # 搜索框
    tk.Label(frame, text="👩‍💼 View Schedule by Name", font=bold_font, bg="white").grid(row=0, column=2, padx=(20, 5), sticky='w')
    name_var = tk.StringVar()
    name_combo = ttk.Combobox(frame, textvariable=name_var, values=sorted(employees), width=10)
    name_combo.grid(row=0, column=3, sticky='w')

    def search_by_name():
        name = name_var.get()
        cur.execute("SELECT date, role, shift FROM schedule WHERE name=? ORDER BY date", (name,))
        rows = cur.fetchall()
        if not rows:
            messagebox.showinfo("No Schedule", f"No schedule found for {name}.")
            return
        result = f"Schedule for {name}:\n"
        for date, role, shift in rows:
            result += f"- {date} ({role}) — {shift} Day\n"
        messagebox.showinfo("Schedule", result)

    tk.Button(frame, text="🔍 Search", command=search_by_name).grid(row=0, column=4, sticky='w', padx=5)

    # 删除记录功能
    def delete_record():
        # 创建删除记录窗口
        delete_window = tk.Toplevel(root)
        delete_window.title("🗑 删除记录")
        delete_window.geometry("400x300")
        delete_window.configure(bg="white")
        
        tk.Label(delete_window, text="🗑 删除排班记录", font=bold_font, bg="white").pack(pady=10)
        
        # 员工选择
        tk.Label(delete_window, text="员工姓名:", font=bold_font, bg="white").pack(anchor="w", padx=20)
        delete_name_var = tk.StringVar()
        delete_name_combo = ttk.Combobox(delete_window, textvariable=delete_name_var, values=sorted(employees), width=20)
        delete_name_combo.pack(pady=5)
        
        # 日期选择
        tk.Label(delete_window, text="日期:", font=bold_font, bg="white").pack(anchor="w", padx=20)
        delete_date_var = tk.StringVar()
        DateEntry(delete_window, textvariable=delete_date_var, date_pattern='yyyy-mm-dd', width=20).pack(pady=5)
        
        # 职位选择
        tk.Label(delete_window, text="职位:", font=bold_font, bg="white").pack(anchor="w", padx=20)
        delete_role_var = tk.StringVar()
        role_combo = ttk.Combobox(delete_window, textvariable=delete_role_var, values=["cleaner", "receptionist"], width=20)
        role_combo.pack(pady=5)
        
        def confirm_delete():
            name = delete_name_var.get()
            date = delete_date_var.get()
            role = delete_role_var.get()
            
            if not name or not date or not role:
                messagebox.showwarning("警告", "请填写完整信息")
                return
            
            # 检查记录是否存在
            cur.execute("SELECT * FROM schedule WHERE name=? AND date=? AND role=?", (name, date, role))
            if not cur.fetchone():
                messagebox.showwarning("警告", "未找到匹配的记录")
                return
            
            if messagebox.askyesno("确认删除", f"确定要删除 {name} 在 {date} 的 {role} 记录吗？"):
                cur.execute("DELETE FROM schedule WHERE name=? AND date=? AND role=?", (name, date, role))
                conn.commit()
                messagebox.showinfo("成功", "记录已删除")
                delete_window.destroy()
        
        tk.Button(delete_window, text="🗑 确认删除", font=bold_font, bg="#e74c3c", fg="white", 
                 command=confirm_delete).pack(pady=20)
    


    # 分配工作
    tk.Label(root, text="🧹 Assign Work", font=bold_font, bg="white").grid(row=2, column=1, sticky='w', padx=30)

    role_frame = tk.Frame(root, bg="white")
    role_frame.grid(row=3, column=1, sticky='w', padx=30)

    roles = ["cleaner", "receptionist"]
    role_labels = ["Cleaner (max 2):", "Receptionist (max 2):"]
    for i, role in enumerate(roles):
        tk.Label(role_frame, text=role_labels[i], bg="white").grid(row=i, column=0)
        emp_var = tk.StringVar()
        emp_combo = ttk.Combobox(role_frame, textvariable=emp_var, values=employees, width=10)
        emp_combo.grid(row=i, column=1)

        shift_var = tk.StringVar(value="Full")
        tk.Radiobutton(role_frame, text="Full Day", variable=shift_var, value="Full", bg="white").grid(row=i, column=2)
        tk.Radiobutton(role_frame, text="Half Day", variable=shift_var, value="Half", bg="white").grid(row=i, column=3)

        def add_work(r=role, e_var=emp_var, s_var=shift_var):
            date = cal.selection_get().strftime("%Y-%m-%d")
            name = e_var.get()
            shift = s_var.get()
            cur.execute("SELECT COUNT(*) FROM schedule WHERE date=? AND role=?", (date, r))
            count = cur.fetchone()[0]
            if count >= 2:
                messagebox.showwarning("Limit", f"Max 2 {r}s per day.")
                return
            cur.execute("INSERT INTO schedule (date, name, role, shift) VALUES (?, ?, ?, ?)", (date, name, r, shift))
            conn.commit()

        def remove_work(r=role, e_var=emp_var):
            date = cal.selection_get().strftime("%Y-%m-%d")
            name = e_var.get()
            cur.execute("DELETE FROM schedule WHERE date=? AND name=? AND role=?", (date, name, r))
            conn.commit()

        tk.Button(role_frame, text="➕ Add", command=add_work).grid(row=i, column=4, padx=5)
        tk.Button(role_frame, text="❌ Remove", command=remove_work).grid(row=i, column=5)

    # 导出部分
    export_frame = tk.LabelFrame(root, text="📦 Export Schedule", font=bold_font, bg="white")
    export_frame.grid(row=3, column=0, padx=10, pady=5, sticky='w')

    start_var = tk.StringVar()
    end_var = tk.StringVar()
    DateEntry(export_frame, textvariable=start_var, date_pattern='yyyy-mm-dd', width=12).grid(row=0, column=1)
    DateEntry(export_frame, textvariable=end_var, date_pattern='yyyy-mm-dd', width=12).grid(row=1, column=1)
    tk.Label(export_frame, text="Start Date:").grid(row=0, column=0, sticky='w')
    tk.Label(export_frame, text="End Date:").grid(row=1, column=0, sticky='w')
    tk.Button(export_frame, text="📁 Export This Period", command=lambda: export_schedule(start_var.get(), end_var.get())).grid(row=2, columnspan=2, pady=5)
    
    # 工资预览功能
    def preview_salary():
        start = start_var.get()
        end = end_var.get()
        
        if not start or not end:
            messagebox.showwarning("警告", "请选择开始和结束日期")
            return
        
        cur.execute("SELECT * FROM schedule WHERE date BETWEEN ? AND ? ORDER BY date, role, name", (start, end))
        rows = cur.fetchall()
        
        if not rows:
            messagebox.showinfo("无数据", "所选时间段内没有排班数据")
            return
        
        # 准备数据
        data = []
        zero_pay_records = []
        for row in rows:
            pay = 0
            if row[3] == "cleaner":
                pay = CLEANER_PAY.get()
            elif row[3] == "receptionist":
                pay = RECEPTIONIST_FULL if row[4] == "Full" else RECEPTIONIST_HALF
            
            data.append([row[1], row[2], row[3], row[4], pay])
            if pay == 0:
                zero_pay_records.append(f"{row[1]} - {row[2]} ({row[3]}) - {row[4]}")
        
        # 创建预览窗口
        preview_window = tk.Toplevel(root)
        preview_window.title("💰 工资预览")
        preview_window.geometry("600x500")
        preview_window.configure(bg="white")
        
        tk.Label(preview_window, text="💰 工资预览", font=bold_font, bg="white", fg="black").pack(pady=10)
        
        # 显示零工资记录警告
        if zero_pay_records:
            warning_frame = tk.Frame(preview_window, bg="#fff3cd", relief="solid", bd=1)
            warning_frame.pack(fill="x", padx=20, pady=10)
            tk.Label(warning_frame, text="⚠️ 发现工资为0的记录:", font=bold_font, bg="#fff3cd", fg="#856404").pack(anchor="w", padx=10, pady=5)
            for record in zero_pay_records:
                tk.Label(warning_frame, text=f"  • {record}", bg="#fff3cd", fg="#856404").pack(anchor="w", padx=20)
        
        # 显示统计信息
        df_preview = pd.DataFrame(data, columns=["Date", "Name", "Role", "Shift", "Pay"])
        total_pay = df_preview['Pay'].sum()
        total_records = len(df_preview)
        
        stats_frame = tk.Frame(preview_window, bg="white")
        stats_frame.pack(fill="x", padx=20, pady=10)
        tk.Label(stats_frame, text=f"总记录数: {total_records}", font=bold_font, bg="white", fg="black").pack(anchor="w")
        tk.Label(stats_frame, text=f"总工资: ${total_pay}", font=bold_font, bg="white", fg="black").pack(anchor="w")
        
        # 显示详细数据
        text_frame = tk.Frame(preview_window, bg="white")
        text_frame.pack(fill="both", expand=True, padx=20, pady=10)
        
        text_widget = tk.Text(text_frame, height=15, width=70)
        scrollbar = tk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_widget.configure(yscrollcommand=scrollbar.set)
        
        text_widget.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # 插入数据
        text_widget.insert("1.0", "详细记录:\n\n")
        for date, name, role, shift, pay in data:
            text_widget.insert("end", f"{date} - {name} ({role}) - {shift} - ${pay}\n")
    
    tk.Button(export_frame, text="💰 工资预览", font=bold_font, bg="#f39c12", fg="white", 
             command=preview_salary).grid(row=3, columnspan=2, pady=5)
    
    # 数据库导入功能
    def import_data():
        import_window = tk.Toplevel(root)
        import_window.title("📥 导入数据")
        import_window.geometry("500x400")
        import_window.configure(bg="white")
        
        tk.Label(import_window, text="📥 导入历史数据", font=bold_font, bg="white", fg="black").pack(pady=10)
        
        # 说明文本
        info_text = """导入格式说明：
1. 选择Excel文件（.xlsx格式）
2. 文件应包含以下列：Date, Name, Role, Shift
3. Date格式：YYYY-MM-DD
4. Role：cleaner 或 receptionist
5. Shift：Full 或 Half

⚠️ 重要提示：
导入时会覆盖当天相同员工相同职位的原数据
确保Excel文件包含完整正确的数据"""
        
        tk.Label(import_window, text=info_text, font=("Arial", 10), bg="white", fg="black", justify="left").pack(pady=10, padx=20)
        
        def select_file():
            file_path = filedialog.askopenfilename(
                title="选择Excel文件",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
            )
            if file_path:
                try:
                    # 读取Excel文件
                    df = pd.read_excel(file_path)
                    required_columns = ['Date', 'Name', 'Role', 'Shift']
                    
                    if not all(col in df.columns for col in required_columns):
                        messagebox.showerror("错误", "Excel文件格式不正确，请确保包含Date, Name, Role, Shift列")
                        return
                    
                    # 导入数据 - 覆盖当天原数据
                    imported_count = 0
                    for _, row in df.iterrows():
                        try:
                            date = str(row['Date'])
                            name = row['Name']
                            role = row['Role']
                            shift = row['Shift']
                            
                            # 先删除当天该员工该职位的原数据
                            cur.execute("DELETE FROM schedule WHERE date=? AND name=? AND role=?", 
                                      (date, name, role))
                            
                            # 插入新数据
                            cur.execute("INSERT INTO schedule (date, name, role, shift) VALUES (?, ?, ?, ?)",
                                      (date, name, role, shift))
                            imported_count += 1
                        except Exception as e:
                            print(f"导入记录失败: {row}, 错误: {e}")
                    
                    conn.commit()
                    messagebox.showinfo("成功", f"成功导入 {imported_count} 条记录\n注意：已覆盖当天相同员工相同职位的原数据")
                    import_window.destroy()
                    
                except Exception as e:
                    messagebox.showerror("错误", f"导入失败: {str(e)}")
        
        tk.Button(import_window, text="📁 选择文件", font=bold_font, bg="#3498db", fg="white",
                 command=select_file).pack(pady=20)
    
    tk.Button(export_frame, text="📥 导入数据", font=bold_font, bg="#27ae60", fg="white",
             command=import_data).grid(row=4, columnspan=2, pady=5)
    
    # 添加房子图片到空白处
    try:
        house_img = Image.open("house.png")
        house_img = house_img.resize((80, 80))
        house_photo = ImageTk.PhotoImage(house_img)
        house_label = tk.Label(root, image=house_photo, bg="white")
        house_label.image = house_photo
        house_label.grid(row=4, column=0, padx=20, pady=20, sticky='sw')
    except:
        # 如果没有图片文件，使用emoji
        house_label = tk.Label(root, text="🏠", font=("Arial", 40), bg="white", fg="#2c3e50")
        house_label.grid(row=4, column=0, padx=20, pady=20, sticky='sw')

build_scheduler_controls(root, employees, cur, conn, CLEANER_PAY, RECEPTIONIST_FULL, RECEPTIONIST_HALF, export_schedule)

root.mainloop()